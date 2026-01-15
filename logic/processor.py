import os
import re
import pandas as pd
from docx import Document
from docxtpl import DocxTemplate
from io import BytesIO
from .rules import get_civility_rules


class MailGenerator:
    def __init__(self, template_path, excel_path, output_dir):
        self.template_path = template_path
        self.excel_path = excel_path
        self.output_dir = output_dir

    @staticmethod
    def extract_tags(template_path):
        if not os.path.exists(template_path):
            return []
        doc = Document(template_path)
        text = [p.text for p in doc.paragraphs]
        for t in doc.tables:
            for r in t.rows:
                for c in r.cells:
                    text.append(c.text)
        found = re.findall(r"\{\{\s*(\w+)\s*\}\}", " ".join(text))
        auto = ['civ_titre', 'politesse_courte', 'formule_politesse_fin']
        return sorted(list(set(t for t in found if t not in auto)))

    def run_direct(self, df, add_page_breaks, progress_callback):
        """
        Génère les documents directement sans toucher au template original.

        Args:
            df: DataFrame avec les données
            add_page_breaks: True pour ajouter des sauts de page, False sinon
            progress_callback: Fonction de progression
        """
        df.columns = [str(c).strip() for c in df.columns]
        master_doc = None
        total = len(df)

        for index, row in df.iterrows():
            # Charger le template original à chaque fois (ne jamais le modifier)
            tpl = DocxTemplate(self.template_path)
            context = row.to_dict()
            context.update(get_civility_rules(row))

            # Rendre le template avec le contexte
            tpl.render(context)

            # Sauvegarder en mémoire
            temp_stream = BytesIO()
            tpl.save(temp_stream)
            temp_stream.seek(0)
            current_doc = Document(temp_stream)

            if master_doc is None:
                # Premier document devient le master
                master_doc = current_doc
            else:
                # Ajouter un saut de page si demandé
                if add_page_breaks:
                    master_doc.add_page_break()

                # Copier les éléments du document courant vers le master
                for element in current_doc.element.body:
                    # Ne pas copier les propriétés de section (sectPr)
                    if not element.tag.endswith('sectPr'):
                        master_doc.element.body.append(element)

            progress_callback(index + 1, total)

        # Sauvegarder le document final
        out_path = os.path.join(self.output_dir, "Courriers_Finalises.docx")
        master_doc.save(out_path)

        return out_path
